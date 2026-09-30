from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = PROJECT_ROOT / "data" / "raw" / "online_retail_II.xlsx"


def main() -> None:
    if not WORKBOOK_PATH.is_file():
        raise FileNotFoundError(f"Workbook not found: {WORKBOOK_PATH}")

    workbook = load_workbook(
        WORKBOOK_PATH,
        read_only=True,
        data_only=True,
    )

    try:
        print(f"Workbook: {WORKBOOK_PATH.name}")
        print(f"Sheets: {workbook.sheetnames}")

        for sheet in workbook.worksheets:
            print(f"\nSheet: {sheet.title}")

            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)
            print(f"Columns: {headers}")

            print("First three data rows:")
            for _ in range(3):
                row = next(rows, None)
                if row is None:
                    break
                print(row)
    finally:
        workbook.close()


if __name__ == "__main__":
    main()